"""
Bóc tách nội dung thô từ file XLSX — PHÂN BIỆT được dòng text (tiêu đề/mô tả,
chỉ có 1 ô có giá trị) và dòng thuộc bảng (>= 2 ô có giá trị), để tránh dòng
text bị "ăn" theo độ rộng cột của bảng bên dưới (gây nhiễu, tốn token).

Không tự diễn giải Ý NGHĨA nội dung ở bước này — chỉ tổ chức lại cho đúng
hình dạng (text thuần vs bảng), dữ liệu vẫn nguyên văn.
"""
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def _build_merge_value_map(sheet) -> dict[tuple[int, int], str]:
    """
    openpyxl không cho ghi trực tiếp vào các ô nằm trong vùng merge (trừ ô
    trên-cùng-trái) — MergedCell.value là read-only. Nên thay vì ghi đè, xây
    1 map (dòng, cột) -> giá trị của ô gốc, để dùng khi đọc dữ liệu.
    """
    merge_map: dict[tuple[int, int], str] = {}
    for merged_range in sheet.merged_cells.ranges:
        top_value = sheet.cell(row=merged_range.min_row, column=merged_range.min_col).value
        if top_value is None:
            continue
        for r in range(merged_range.min_row, merged_range.max_row + 1):
            for c in range(merged_range.min_col, merged_range.max_col + 1):
                merge_map[(r, c)] = top_value
    return merge_map


def _read_rows_with_position(sheet, merge_map: dict[tuple[int, int], str]) -> list[list[tuple[int, str]]]:
    """Đọc từng dòng, chỉ giữ lại các ô CÓ giá trị (kể cả giá trị suy ra từ
    vùng merge), kèm số cột gốc của ô đó."""
    rows = []
    for row in sheet.iter_rows():
        cells = []
        for cell in row:
            val = cell.value
            if val is None:
                val = merge_map.get((cell.row, cell.column))
            if val is not None and str(val).strip() != "":
                cells.append((cell.column, str(val).strip()))
        rows.append(cells)
    return rows


def _group_rows_into_blocks(rows: list[list[tuple[int, str]]]) -> list[list]:
    """
    Gom các dòng thành block:
    - Dòng rỗng hoàn toàn: kết thúc block bảng đang gom (nếu có), không tạo block.
    - Dòng chỉ có 1 ô: 1 block text riêng (đứng độc lập).
    - Dòng có >= 2 ô: gom vào block bảng đang mở (liên tiếp nhau).
    """
    blocks = []
    current_table: list = []

    def flush_table():
        if current_table:
            blocks.append(("table", current_table.copy()))
            current_table.clear()

    for cells in rows:
        if not cells:
            flush_table()
            continue

        distinct_values = {val for _, val in cells}
        if len(cells) == 1 or len(distinct_values) == 1:
            # 1 ô có giá trị, HOẶC nhiều ô nhưng cùng 1 giá trị (thường do merge
            # nhiều cột làm tiêu đề) -> coi là dòng text, không phải dòng bảng.
            flush_table()
            blocks.append(("text", [[cells[0]]]))
        else:
            current_table.append(cells)
    flush_table()
    return blocks


def _render_text_block(block) -> str:
    return block[0][0][1]


def _render_table_block(block) -> str:
    columns = sorted({col for row in block for col, _ in row})
    if not columns:
        return ""

    header_map = dict(block[0])
    data_rows = block[1:]

    def col_label(c):
        # Ưu tiên dùng chính nội dung dòng đầu làm tên cột (thường là header thật
        # của bảng). Nếu ô đó trống, dùng ký hiệu cột Excel (A, B, C...) làm nhãn.
        return header_map.get(c) or get_column_letter(c)

    lines = [
        "| " + " | ".join(col_label(c) for c in columns) + " |",
        "| " + " | ".join(["---"] * len(columns)) + " |",
    ]
    for row in data_rows:
        row_map = dict(row)
        lines.append("| " + " | ".join(row_map.get(c, "") for c in columns) + " |")
    return "\n".join(lines)


def extract_text(file_path: str) -> str:
    workbook = load_workbook(file_path, data_only=True)
    output_parts = []

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        merge_map = _build_merge_value_map(sheet)

        output_parts.append(f"## Sheet: {sheet_name}")

        rows = _read_rows_with_position(sheet, merge_map)
        blocks = _group_rows_into_blocks(rows)

        for block_type, block in blocks:
            if block_type == "text":
                output_parts.append(_render_text_block(block))
            else:
                rendered = _render_table_block(block)
                if rendered:
                    output_parts.append(rendered)

    return "\n\n".join(output_parts)


# if __name__ == "__main__":
#     import os

#     input_path = os.path.join(os.path.dirname(__file__), "..", "data", "input", "data_llm.xlsx")
#     output_dir = os.path.join(os.path.dirname(__file__), "..", "data", "extract")
#     os.makedirs(output_dir, exist_ok=True)

#     file_name = os.path.splitext(os.path.basename(input_path))[0]
#     output_path = os.path.join(output_dir, f"result_{file_name}.md")

#     content = extract_text(input_path)
#     with open(output_path, "w", encoding="utf-8") as f:
#         f.write(content)

#     print(f"Done: {output_path}")