from __future__ import annotations

import io
from typing import Any

import openpyxl

from core.checklist_models import ChecklistItem


def _clean(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def _header_matches(row: list[Any]) -> bool:
    texts = [(_clean(v).lower()) for v in row if v is not None]
    if not texts:
        return False
    normalized = " ".join(texts)
    return "分類" in normalized and "チェック項目" in normalized and "回答" in normalized


def parse_checklist_workbook(file_like: Any) -> list[ChecklistItem]:
    if isinstance(file_like, bytes):
        blob = file_like
    elif hasattr(file_like, "getvalue"):
        blob = file_like.getvalue()
    else:
        blob = file_like.read()

    workbook = openpyxl.load_workbook(io.BytesIO(blob), data_only=True)

    target_sheet = None
    for ws in workbook.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=min(25, ws.max_row), values_only=True):
            if _header_matches(list(row)):
                target_sheet = ws
                break
        if target_sheet is not None:
            break

    if target_sheet is None:
        raise ValueError("チェックリストシートを特定できませんでした。ヘッダー行を確認してください。")

    rows = list(target_sheet.iter_rows(values_only=True))
    header_index = None
    for idx, row in enumerate(rows):
        if _header_matches(list(row)):
            header_index = idx
            break

    if header_index is None:
        raise ValueError("有効なチェックリストヘッダーが見つかりませんでした。")

    header_row = rows[header_index]
    mapped: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        text = _clean(cell)
        if not text:
            continue
        if "分類" in text:
            mapped["category"] = idx
        if "チェック内容" in text:
            mapped["detail"] = idx
        if "チェック項目" in text:
            mapped["question"] = idx
        if "回答" in text:
            mapped["answer"] = idx
        if "コメント" in text and "調達先コメント" in text:
            mapped["comment"] = idx
        if "調達先コメント" in text:
            mapped["comment"] = idx
        if "是正依頼回答" in text:
            mapped["corrective_response"] = idx
        if text.lower() == "no":
            mapped["no"] = idx

    if "question" not in mapped:
        raise ValueError("チェック項目列が見つかりませんでした。")

    required_columns = {"category", "detail", "question", "answer", "comment"}
    missing_columns = required_columns - mapped.keys()
    if missing_columns:
        raise ValueError(f"必須列が見つかりません: {', '.join(sorted(missing_columns))}")

    items: list[ChecklistItem] = []
    item_id_counts: dict[str, int] = {}
    previous_values: dict[str, str] = {}
    for r_index, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        if row is None:
            continue
        if all(_clean(v) == "" for v in row):
            continue

        category = _clean(row[mapped["category"]] if mapped["category"] < len(row) else "")
        question = _clean(row[mapped["question"]] if mapped["question"] < len(row) else "")
        detail = _clean(row[mapped["detail"]] if mapped["detail"] < len(row) else "")
        answer = _clean(row[mapped["answer"]] if mapped["answer"] < len(row) else "")
        comment = _clean(row[mapped["comment"]] if mapped["comment"] < len(row) else "")
        no_value = _clean(row[mapped.get("no", 0)] if mapped.get("no", 0) < len(row) else "")
        corrective_response = _clean(row[mapped["corrective_response"]] if "corrective_response" in mapped and mapped["corrective_response"] < len(row) else "")

        # openpyxl exposes merged follow-on cells as None; carry only the
        # descriptive fields forward, never partner answers/comments.
        for key, value in (("category", category), ("detail", detail)):
            if value:
                previous_values[key] = value
            elif key in previous_values:
                if key == "category":
                    category = previous_values[key]
                else:
                    detail = previous_values[key]

        if not question and not detail and not answer and not comment:
            continue

        item_id = question or f"row-{r_index}"
        if no_value and question:
            item_id = f"{no_value}-{question.split()[0]}" if question.split() else question
        item_id_counts[item_id] = item_id_counts.get(item_id, 0) + 1
        if item_id_counts[item_id] > 1:
            item_id = f"{item_id}-{item_id_counts[item_id]}"

        item = ChecklistItem(
            item_id=item_id,
            source_row=r_index,
            source_sheet=target_sheet.title,
            source_columns=[openpyxl.utils.get_column_letter(index + 1) for index in range(len(header_row)) if header_row[index] is not None],
            category=category,
            question=question,
            detail=detail,
            partner_answer=answer,
            partner_comment=comment,
            corrective_action_response=corrective_response,
            confirmation_required=False,
            confirmation_reason=None,
            result_classification=None,
            auditor_comment=None,
            source_raw={
                "category": category,
                "question": question,
                "detail": detail,
                "answer": answer,
                "comment": comment,
            },
        )
        items.append(item)

    return items
