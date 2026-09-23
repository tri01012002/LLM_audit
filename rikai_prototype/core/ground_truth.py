from __future__ import annotations

import io
from collections import defaultdict
from typing import Any

import openpyxl

from core.checklist_models import AIAnalysisResult, ChecklistItem, GroundTruthGroup, GroundTruthItem


def _clean(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def _sheet(blob: bytes, name: str):
    workbook = openpyxl.load_workbook(io.BytesIO(blob), data_only=True)
    if name not in workbook.sheetnames:
        raise ValueError(f"Ground-truth sheet not found: {name}")
    return workbook[name]


def load_ground_truth(blob: bytes) -> tuple[dict[str, GroundTruthItem], dict[str, GroundTruthGroup], list[str]]:
    item_sheet = _sheet(blob, "⑤評価根拠_128項目")
    group_sheet = _sheet(blob, "⑥判定サマリー_デモ用")
    policy_sheet = _sheet(blob, "④判定方針_デモ用")

    items: dict[str, GroundTruthItem] = {}
    for row in item_sheet.iter_rows(min_row=2, values_only=True):
        source_row = row[0]
        item_id = _clean(row[2])
        if not item_id or not isinstance(source_row, int):
            continue
        items[item_id] = GroundTruthItem(
            source_row=source_row,
            group_id=_clean(row[1]),
            item_id=item_id,
            answer=_clean(row[3]),
            comment=_clean(row[4]),
            confirmation=_clean(row[5]),
            confirmation_request=_clean(row[6]),
            corrective_action=_clean(row[7]),
            corrective_action_proposal=_clean(row[8]),
            rationale=_clean(row[9]),
            source=[part.strip() for part in _clean(row[10]).split("・") if part.strip()],
            cross_item_consistency=_clean(row[11]),
        )

    groups: dict[str, GroundTruthGroup] = {}
    for row in group_sheet.iter_rows(min_row=22, values_only=True):
        group_id = _clean(row[0])
        if not group_id or not str(group_id).isdigit():
            continue
        groups[group_id] = GroundTruthGroup(
            group_id=group_id,
            category=_clean(row[1]),
            classification=_clean(row[2]),
            corrective_action_count=int(row[3] or 0),
            confirmation_count=int(row[4] or 0),
            summary=_clean(row[5]),
        )

    policy_lines = []
    for row in policy_sheet.iter_rows(min_row=7, max_row=26, values_only=True):
        values = [_clean(value) for value in row[:4] if _clean(value)]
        if values:
            policy_lines.append(" | ".join(values))

    return items, groups, policy_lines


def _item_key(item: ChecklistItem) -> str:
    return item.question.strip()


def aggregate_group_results(items: list[ChecklistItem], results: dict[str, AIAnalysisResult]) -> list[dict[str, Any]]:
    grouped: dict[str, list[ChecklistItem]] = defaultdict(list)
    for item in items:
        grouped[item.group_id or item.item_id.split("-", 1)[0]].append(item)

    summaries = []
    for group_id, group_items in sorted(grouped.items(), key=lambda pair: int(pair[0]) if pair[0].isdigit() else pair[0]):
        group_results = [results[item.item_id] for item in group_items if item.item_id in results]
        confirmations = sum(result.confirmation_required for result in group_results)
        corrective = sum(bool(result.issue_or_risk or result.improvement_proposal) for result in group_results)
        if any(result.status == "AI_ERROR" for result in group_results):
            classification = "AI_ERROR"
        elif any(result.status == "POSSIBLE_CONTRADICTION" for result in group_results):
            classification = "要確認（仮）"
        elif any(result.status in {"NEEDS_INFORMATION", "NEEDS_CONFIRMATION", "NEEDS_REVIEW"} for result in group_results):
            classification = "是正案あり（仮）"
        else:
            classification = "〇（仮）"
        summaries.append({
            "group_id": group_id,
            "category": group_items[0].category,
            "classification": classification,
            "corrective_action_count": corrective,
            "confirmation_count": confirmations,
            "summary": f"{len(group_items)} items; {confirmations} confirmation(s); {corrective} proposed action(s). Demo draft only; evidence not independently verified.",
            "item_ids": [item.item_id for item in group_items],
        })
    return summaries


def compare_ground_truth(items: list[ChecklistItem], results: dict[str, AIAnalysisResult], ground_truth_items: dict[str, GroundTruthItem], ground_truth_groups: dict[str, GroundTruthGroup]) -> dict[str, Any]:
    item_metrics = {"total": 0, "confirmation_matched": 0, "corrective_action_matched": 0}
    for item in items:
        expected = ground_truth_items.get(_item_key(item))
        result = results.get(item.item_id)
        if not expected or not result:
            continue
        item_metrics["total"] += 1
        if ("要" if result.confirmation_required else "不要") == expected.confirmation:
            item_metrics["confirmation_matched"] += 1
        proposed = bool(result.issue_or_risk or result.improvement_proposal)
        if ("有" if proposed else "無") == expected.corrective_action:
            item_metrics["corrective_action_matched"] += 1

    generated_groups = {group["group_id"]: group for group in aggregate_group_results(items, results)}
    group_metrics = {"total": len(ground_truth_groups), "classification_matched": 0}
    for group_id, expected in ground_truth_groups.items():
        actual = generated_groups.get(group_id)
        if actual and actual["classification"] == expected.classification:
            group_metrics["classification_matched"] += 1
    return {"items": item_metrics, "groups": group_metrics}
