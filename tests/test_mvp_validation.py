from __future__ import annotations

import os
import sys
import tempfile
import unittest

from openpyxl import load_workbook

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
APP_ROOT = os.path.join(PROJECT_ROOT, "rikai_prototype")
if APP_ROOT not in sys.path:
    sys.path.insert(0, APP_ROOT)

from core.ai_analyzer import analyze_checklist_items, export_reviewed_workbook
from core.business_rules import evaluate_business_rules
from core.checklist_models import ChecklistItem
from core.excel_parser import parse_checklist_workbook


SAMPLE = os.path.join(
    PROJECT_ROOT,
    "data_test_llm",
    "（送付物）IT調達先チェックリスト_アスノシステム株式会社.xlsx",
)


class MvpValidationTests(unittest.TestCase):
    def test_parser_and_demo_batch(self):
        source = open(SAMPLE, "rb").read()
        items = parse_checklist_workbook(source)
        results = analyze_checklist_items(items, mode="demo")
        self.assertEqual(len(items), 128)
        self.assertEqual(len(results), 128)
        self.assertEqual(items[0].source_row, 2)
        self.assertEqual(items[-1].source_row, 129)
        self.assertTrue(all(item.partner_answer is not None for item in items))
        self.assertTrue(all(result.analysis_method == "DEMO" for result in results.values()))

    def test_rule_cases(self):
        cases = [
            ("〇", "Annual review is performed.", "NORMAL"),
            ("〇", "Currently under review and has not yet been implemented.", "POSSIBLE_CONTRADICTION"),
            ("〇", "Planned to implement next quarter.", "POSSIBLE_CONTRADICTION"),
            ("△", "Currently reviewing the implementation.", "NEEDS_CONFIRMATION"),
            ("✕", "Not implemented yet.", "NEEDS_REVIEW"),
            ("－", "Not applicable to this service.", "SUPPORTED_NA"),
            ("－", "", "NEEDS_CONFIRMATION"),
            ("", "", "NEEDS_INFORMATION"),
            ("〇", "Yes, done.", "NEEDS_REVIEW"),
        ]
        for answer, comment, expected in cases:
            item = ChecklistItem(
                item_id="case",
                source_row=2,
                source_sheet="Checklist",
                question="Requirement",
                detail="Detail",
                partner_answer=answer,
                partner_comment=comment,
            )
            self.assertEqual(evaluate_business_rules(item)["status"], expected)

    def test_export_preserves_partner_k_and_writes_auditor_fields(self):
        source = open(SAMPLE, "rb").read()
        items = parse_checklist_workbook(source)
        results = analyze_checklist_items(items, mode="demo")
        original = load_workbook(SAMPLE, data_only=False)["③チェックリスト"]
        decisions = {
            item_id: {
                "final_confirmation_request": result.confirmation_reason or "",
                "final_proposal": result.improvement_proposal or "",
                "final_comment": "Auditor-approved comment",
                "confirmation_required": result.confirmation_required,
                "issue_or_risk": result.issue_or_risk or "",
                "classification": result.status,
            }
            for item_id, result in results.items()
        }
        output = os.path.join(tempfile.gettempdir(), "rai_mvp_validation_test.xlsx")
        try:
            export_reviewed_workbook(source, items, decisions, output)
            exported = load_workbook(output, data_only=False)["③チェックリスト"]
            self.assertEqual(exported["K2"].value, original["K2"].value)
            self.assertEqual(exported["L2"].value, results[items[0].item_id].status)
            self.assertEqual(exported["M2"].value, "Auditor-approved comment")
            self.assertEqual(exported["L3"].value, results[items[1].item_id].status)
        finally:
            if os.path.exists(output):
                os.remove(output)


if __name__ == "__main__":
    unittest.main()
